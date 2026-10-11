"""
Report builder: turns the analysis done in a session into a Word (.docx) and PDF report.

The report is first described as simple "blocks" (title, headings, paragraphs, bullets, tables,
images) and then rendered to both formats, so the two always match.
"""
import datetime
import io
import os
import re
from xml.sax.saxutils import escape

import pandas as pd

# ----------------------------------------------------------------------------- text helpers
INLINE = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`")


def inline_runs(text):
    """Split markdown inline text into (text, bold, italic, code) runs."""
    runs, pos = [], 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            runs.append((text[pos:m.start()], False, False, False))
        if m.group(1) is not None:
            runs.append((m.group(1), True, False, False))
        elif m.group(2) is not None:
            runs.append((m.group(2), False, True, False))
        else:
            runs.append((m.group(3), False, False, True))
        pos = m.end()
    if pos < len(text):
        runs.append((text[pos:], False, False, False))
    return runs or [("", False, False, False)]


def _fmt(v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
        return ""
    if isinstance(v, (pd.Timestamp, datetime.datetime, datetime.date)):
        return pd.Timestamp(v).strftime("%Y-%m-%d")
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int) or (hasattr(v, "dtype") and "int" in str(getattr(v, "dtype", ""))):
        return f"{int(v):,}"
    if isinstance(v, float) or "float" in str(getattr(v, "dtype", "")):
        v = float(v)
        return f"{int(v):,}" if v.is_integer() and abs(v) < 1e15 else f"{v:,.3f}".rstrip("0").rstrip(".")
    return str(v)


def df_block(caption, df, max_rows=25, max_cols=9):
    cols = list(df.columns)[:max_cols]
    d = df[cols].head(max_rows)
    note = []
    if len(df) > max_rows:
        note.append(f"first {max_rows} of {len(df):,} rows")
    if df.shape[1] > max_cols:
        note.append(f"first {max_cols} of {df.shape[1]} columns")
    cap = caption + (f" ({', '.join(note)}; full table is in the Excel download)" if note else "")
    return ("table", [str(c) for c in cols], [[_fmt(v) for v in row] for row in d.itertuples(index=False, name=None)], cap)


def md_to_blocks(md):
    """Convert the AI's markdown narrative (headings, paragraphs, lists, tables) into blocks."""
    blocks, para, bullets, nums, tbl = [], [], [], [], []

    def flush():
        nonlocal para, bullets, nums, tbl
        if para:
            blocks.append(("p", " ".join(para)))
        if bullets:
            blocks.append(("bullets", bullets))
        if nums:
            blocks.append(("numbered", nums))
        if tbl:
            rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in tbl
                    if not re.fullmatch(r"\s*\|?[\s:\-|]+\|?\s*", r)]
            if rows:
                w = len(rows[0])
                blocks.append(("table", rows[0], [(r + [""] * w)[:w] for r in rows[1:]], ""))
        para, bullets, nums, tbl = [], [], [], []

    for line in md.splitlines():
        s = line.strip()
        if not s:
            flush()
        elif s.startswith("|"):
            if para or bullets or nums:
                flush()
            tbl.append(s)
        elif re.match(r"^#{1,4}\s", s):
            flush()
            lvl = len(s) - len(s.lstrip("#"))
            blocks.append(("h", min(max(lvl, 1), 3), s.lstrip("#").strip()))
        elif re.match(r"^[-*•]\s+", s):
            if para or nums or tbl:
                flush()
            bullets.append(re.sub(r"^[-*•]\s+", "", s))
        elif re.match(r"^\d+[.)]\s+", s):
            if para or bullets or tbl:
                flush()
            nums.append(re.sub(r"^\d+[.)]\s+", "", s))
        else:
            if bullets or nums or tbl:
                flush()
            para.append(s)
    flush()
    return blocks


# ----------------------------------------------------------------------------- content
def context_text(ws, messages):
    """Facts handed to the AI so it can write the report narrative (no invented numbers)."""
    parts = [ws.profile()[:20000]] if ws.frames else []
    for key, lines in ws.clean_log.items():
        txt = "\n".join(l for l in lines if l)
        if txt:
            parts.append(f"CLEANING APPLIED TO {key}:\n{txt}")
    if ws.formulas_used:
        parts.append("EXCEL FORMULAS USED:\n" + "\n".join(f"{k}: {f} -> {r}" for k, f, r in ws.formulas_used[-25:]))
    for title, t in ws.hist_tables[-8:]:
        parts.append(f"RESULT TABLE: {title}\n{t.head(12).round(3).to_csv(index=False)}")
    if messages:
        convo = "\n".join(f"{m['role'].upper()}: {m['content'][:1500]}" for m in messages[-12:])
        parts.append("CONVERSATION SO FAR:\n" + convo)
    return "\n\n".join(parts)[:60000]


def build_blocks(ws, title, subtitle, ai_md="", include_overview=True, include_tables=True,
                 include_charts=True, include_formulas=True):
    B = [("title", title, subtitle)]
    if ai_md.strip():
        B += [("h", max(1, b[1] - 1), b[2]) if b[0] == "h" else b for b in md_to_blocks(ai_md)]
    if include_overview and ws.frames:
        B.append(("h", 1, "Data overview"))
        for key, df in ws.frames.items():
            B.append(("h", 2, key))
            B.append(("p", f"{len(df):,} rows and {df.shape[1]} columns."))
            cls = ws._class_table(df.iloc[:, :30])[["excel_col", "column", "variable_type", "measurement_scale", "unique_values", "missing"]]
            B.append(df_block("Variable classification", cls, max_rows=30))
    if ws.frames:
        B.append(("h", 1, "Data cleaning"))
        logs = {k: [l for l in v if l] for k, v in ws.clean_log.items() if any(v)}
        if logs:
            for key, lines in logs.items():
                B.append(("h", 2, key))
                B.append(("bullets", lines))
        else:
            B.append(("p", "No data cleaning was applied. The analysis used the data exactly as uploaded."))
    tables = [(t, d) for t, d in ws.hist_tables if not t.startswith("Variable classification")]
    if include_tables and tables:
        B.append(("h", 1, "Analysis results"))
        for t, d in tables[-12:]:
            B.append(df_block(t, d))
    if include_charts and ws.hist_figs:
        B.append(("h", 1, "Charts"))
        for t, png in ws.hist_figs:
            B.append(("image", png, t))
    if include_formulas and ws.formulas_used:
        B.append(("h", 1, "Excel formulas used"))
        B.append(("table", ["Dataset", "Formula", "Result"],
                  [[k.split(" / ")[-1] or k, f, r] for k, f, r in ws.formulas_used[-40:]], "Formulas evaluated on the uploaded data"))
    B.append(("note", "All figures in this report were calculated by code from the uploaded data. Narrative text "
                      "(summary, findings, recommendations) is AI-written and should be reviewed before it is relied on."))
    return B


# ----------------------------------------------------------------------------- DOCX
def to_docx(blocks):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.8)
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)

    def runs(par, text, size=None, italic=False):
        for t, b, i, c in inline_runs(text):
            r = par.add_run(t)
            r.bold, r.italic = b, (i or italic)
            if c:
                r.font.name = "Consolas"
            if size:
                r.font.size = Pt(size)

    for blk in blocks:
        k = blk[0]
        if k == "title":
            d.add_heading(blk[1], 0)
            p = d.add_paragraph()
            runs(p, blk[2], italic=True)
        elif k == "h":
            d.add_heading(blk[2], level=blk[1])
        elif k == "p":
            runs(d.add_paragraph(), blk[1])
        elif k == "bullets":
            for it in blk[1]:
                runs(d.add_paragraph(style="List Bullet"), it)
        elif k == "numbered":
            for n, it in enumerate(blk[1], 1):
                p = d.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.first_line_indent = Inches(-0.2)
                runs(p, f"{n}. {it}")
        elif k == "table":
            _, head, rows, cap = blk
            if cap:
                runs(d.add_paragraph(), f"**{cap}**", size=9)
            t = d.add_table(rows=1, cols=len(head))
            t.style = "Light Grid Accent 1"
            size = 9 if len(head) <= 5 else 8 if len(head) <= 7 else 7
            for i, h in enumerate(head):
                cell = t.rows[0].cells[i]
                cell.text = ""
                runs(cell.paragraphs[0], f"**{h}**", size=size)
            for row in rows:
                cells = t.add_row().cells
                for i, v in enumerate(row[:len(head)]):
                    cells[i].text = ""
                    runs(cells[i].paragraphs[0], v, size=size)
            d.add_paragraph()
        elif k == "image":
            d.add_picture(io.BytesIO(blk[1]), width=Inches(6))
            d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            runs(p, blk[2], size=9, italic=True)
        elif k == "note":
            p = d.add_paragraph()
            runs(p, blk[1], size=9, italic=True)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    bio = io.BytesIO()
    d.save(bio)
    return bio.getvalue()


# ----------------------------------------------------------------------------- PDF
_MAP = {"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"', "\u2013": "-", "\u2014": "-", "\u2022": "-",
        "\u2026": "...", "\u2192": "->", "\u2190": "<-", "\u2265": ">=", "\u2264": "<=", "\u2248": "~", "\u00a0": " ",
        "\u2713": "v", "\u2714": "v", "\u00d7": "x"}


def _pdf_text(t):
    for a, b in _MAP.items():
        t = t.replace(a, b)
    return t.encode("latin-1", "replace").decode("latin-1")


def to_pdf(blocks):
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    base, bold, ital, mono = "Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Courier"
    fdir = os.path.join(os.path.dirname(reportlab.__file__), "fonts")
    try:
        for name, f in (("Vera", "Vera.ttf"), ("Vera-B", "VeraBd.ttf"), ("Vera-I", "VeraIt.ttf")):
            pdfmetrics.registerFont(TTFont(name, os.path.join(fdir, f)))
        pdfmetrics.registerFontFamily("Vera", normal="Vera", bold="Vera-B", italic="Vera-I", boldItalic="Vera-B")
        base, bold, ital = "Vera", "Vera-B", "Vera-I"
    except Exception:
        pass

    ss = getSampleStyleSheet()
    S = {
        "title": ParagraphStyle("t", parent=ss["Title"], fontName=bold, fontSize=22, leading=26, alignment=0, spaceAfter=6),
        "sub": ParagraphStyle("s", parent=ss["Normal"], fontName=ital, fontSize=10, textColor=colors.grey, spaceAfter=12),
        "h1": ParagraphStyle("h1", parent=ss["Heading1"], fontName=bold, fontSize=16, spaceBefore=14, spaceAfter=6, textColor=colors.HexColor("#c2410c")),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontName=bold, fontSize=13, spaceBefore=10, spaceAfter=4),
        "h3": ParagraphStyle("h3", parent=ss["Heading3"], fontName=bold, fontSize=11, spaceBefore=8, spaceAfter=3),
        "p": ParagraphStyle("p", parent=ss["Normal"], fontName=base, fontSize=10, leading=14, spaceAfter=6),
        "li": ParagraphStyle("li", parent=ss["Normal"], fontName=base, fontSize=10, leading=14, leftIndent=16, bulletIndent=4, spaceAfter=2),
        "cell": ParagraphStyle("c", parent=ss["Normal"], fontName=base, fontSize=7.5, leading=9.5),
        "cellh": ParagraphStyle("ch", parent=ss["Normal"], fontName=bold, fontSize=7.5, leading=9.5, textColor=colors.white),
        "cap": ParagraphStyle("cap", parent=ss["Normal"], fontName=ital, fontSize=8.5, textColor=colors.HexColor("#444444"), spaceAfter=4, spaceBefore=4),
        "note": ParagraphStyle("n", parent=ss["Normal"], fontName=ital, fontSize=8, textColor=colors.grey, spaceBefore=10),
    }

    def x(text):
        out = []
        for t, b, i, c in inline_runs(_pdf_text(text)):
            t = escape(t)
            if c:
                t = f'<font face="{mono}">{t}</font>'
            if b:
                t = f"<b>{t}</b>"
            if i:
                t = f"<i>{t}</i>"
            out.append(t)
        return "".join(out)

    W = 17 * cm
    story = []
    for blk in blocks:
        k = blk[0]
        if k == "title":
            story += [Paragraph(x(blk[1]), S["title"]), Paragraph(x(blk[2]), S["sub"])]
        elif k == "h":
            story.append(Paragraph(x(blk[2]), S[f"h{blk[1]}"]))
        elif k == "p":
            story.append(Paragraph(x(blk[1]), S["p"]))
        elif k == "bullets":
            story += [Paragraph(x(it), S["li"], bulletText="-") for it in blk[1]]
            story.append(Spacer(1, 4))
        elif k == "numbered":
            story += [Paragraph(x(it), S["li"], bulletText=f"{n}.") for n, it in enumerate(blk[1], 1)]
            story.append(Spacer(1, 4))
        elif k == "table":
            _, head, rows, cap = blk
            data = [[Paragraph(x(h), S["cellh"]) for h in head]] + [[Paragraph(x(v), S["cell"]) for v in r[:len(head)]] for r in rows]
            t = Table(data, colWidths=[W / len(head)] * len(head), repeatRows=1)
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8590c")),
                                   ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cccccc")),
                                   ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#faf3ee")]),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]))
            story += ([Paragraph(x(cap), S["cap"])] if cap else []) + [t, Spacer(1, 8)]
        elif k == "image":
            w, h = ImageReader(io.BytesIO(blk[1])).getSize()
            iw = min(W, 15 * cm)
            story.append(KeepTogether([Image(io.BytesIO(blk[1]), width=iw, height=iw * h / w), Paragraph(x(blk[2]), S["cap"]), Spacer(1, 8)]))
        elif k == "note":
            story.append(Paragraph(x(blk[1]), S["note"]))
    bio = io.BytesIO()
    SimpleDocTemplate(bio, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm,
                      title="Report").build(story)
    return bio.getvalue()
